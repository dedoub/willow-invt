"""급여명세서를 워드 문서로 짠다.

한컴·엑셀 서식은 눈으로만 참고했다. 그 파일들은 굴림체·맑은 고딕을 쓰는데 이 맥에
없어서 PDF 로 뽑으면 한글이 통째로 빠진다. 그래서 문서를 새로 쓰고 글꼴은 여기 있는
것으로 고정한다.
"""
import datetime
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = 'Nanum Gothic'
HEAD_FILL = 'DCE6F1'        # 표 머리·항목 칸의 옅은 파랑
LINE = 'BFBFBF'

PAY_LABELS = ['기 본 급', '식    대', '차량유지비', '보육수당', '근속수당', '직책수당',
              '연장수당', '가족수당', '상 여 금', '기    타', '', '']
DEDUCT_LABELS = ['국민연금', '건강보험', '장기요양보험', '고용보험', '소 득 세', '지방소득세',
                 '상조회비', '가 불 금', '기    타', '', '', '']


def _font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)


def _shade(cell, fill):
    element = OxmlElement('w:shd')
    element.set(qn('w:val'), 'clear')
    element.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(element)


def _borders(table):
    properties = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), LINE)
        borders.append(element)
    properties.append(borders)


def _write(cell, text, *, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, fill=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    _font(paragraph.add_run(text), size=size, bold=bold)
    if fill:
        _shade(cell, fill)


def _spacer(document, points=6):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(points)
    _font(paragraph.add_run(''), size=4)


def _heading(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    _font(paragraph.add_run(f'■ {text}'), size=10, bold=True)


def _won(amount):
    # 항목이 없는 칸은 비운다. 서식의 줄은 고정이라 안 쓰는 항목이 늘 남는데, 거기에 '원'
    # 만 찍히면 0원을 지급하거나 공제한 것처럼 읽힌다(CEO 2026-09-22).
    return f'{amount:,}원' if amount else ''


def build_payslip(*, path, name, department, title, period, paid_on,
                  pays, deducts, pay_total, deduct_total, net, net_korean,
                  company, ceo_title, seal=None):
    """명세서 한 장. pays/deducts 는 [(이름, 금액)] 이고 서식의 줄 순서를 따른다."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = document.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(14)
    _font(heading.add_run('급 여 명 세 서'), size=20, bold=True, color=RGBColor(0x1F, 0x49, 0x7D))

    who = document.add_table(rows=1, cols=6)
    _borders(who)
    widths = [Cm(2.0), Cm(4.0), Cm(2.0), Cm(3.0), Cm(2.0), Cm(3.0)]
    labels = [('부    서', department), ('직    책', title), ('성    명', name)]
    cells = who.rows[0].cells
    for index, (label, value) in enumerate(labels):
        _write(cells[index * 2], label, fill=HEAD_FILL)
        _write(cells[index * 2 + 1], value)
    for index, width in enumerate(widths):
        for cell in who.columns[index].cells:
            cell.width = width

    _heading(document, '지급 및 공제내역')

    rows = max(len(PAY_LABELS), len(DEDUCT_LABELS))
    table = document.add_table(rows=rows + 2, cols=4)
    _borders(table)
    head = table.rows[0].cells
    head[0].merge(head[1])
    head = table.rows[0].cells
    _write(head[0], '지 급 항 목', bold=True, fill=HEAD_FILL)
    merged = table.rows[0].cells[2].merge(table.rows[0].cells[3])
    _write(merged, '공 제 항 목', bold=True, fill=HEAD_FILL)

    pay_map = dict(pays)
    deduct_map = dict(deducts)
    for index in range(rows):
        cells = table.rows[index + 1].cells
        pay_label = PAY_LABELS[index] if index < len(PAY_LABELS) else ''
        deduct_label = DEDUCT_LABELS[index] if index < len(DEDUCT_LABELS) else ''
        _write(cells[0], pay_label, fill=HEAD_FILL)
        _write(cells[1], _won(pay_map.get(pay_label.replace(' ', ''), 0)), align=WD_ALIGN_PARAGRAPH.RIGHT)
        _write(cells[2], deduct_label, fill=HEAD_FILL)
        _write(cells[3], _won(deduct_map.get(deduct_label.replace(' ', ''), 0)), align=WD_ALIGN_PARAGRAPH.RIGHT)

    last = table.rows[rows + 1].cells
    _write(last[0], '지급합계', bold=True, fill=HEAD_FILL)
    _write(last[1], _won(pay_total), bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _write(last[2], '공제합계', bold=True, fill=HEAD_FILL)
    _write(last[3], _won(deduct_total), bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    for index, width in enumerate([Cm(3.4), Cm(4.6), Cm(3.4), Cm(4.6)]):
        for cell in table.columns[index].cells:
            cell.width = width

    _heading(document, '종합내역')

    summary = document.add_table(rows=4, cols=2)
    _borders(summary)
    lines = [
        ('기본근무 산출근거', period),
        ('총    액', f'{pay_total:,}원'),
        ('공제총액', f'{deduct_total:,}원'),
        ('실수령액', f'일금  {net_korean}원정  (\\ {net:,})'),
    ]
    for index, (label, value) in enumerate(lines):
        cells = summary.rows[index].cells
        _write(cells[0], label, bold=True, fill=HEAD_FILL)
        _write(cells[1], value)
        cells[0].width = Cm(3.4)
        cells[1].width = Cm(12.6)

    _spacer(document, 14)
    thanks = document.add_paragraph()
    thanks.alignment = WD_ALIGN_PARAGRAPH.CENTER
    thanks.paragraph_format.space_after = Pt(16)
    _font(thanks.add_run('귀하의 노고에 감사 드립니다.'), size=10)

    when = document.add_paragraph()
    when.alignment = WD_ALIGN_PARAGRAPH.CENTER
    when.paragraph_format.space_after = Pt(22)
    _font(when.add_run(f'{paid_on.year}년 {paid_on.month:02d}월 {paid_on.day:02d}일'), size=10)

    signer = document.add_paragraph()
    signer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    signer.paragraph_format.space_after = Pt(0)
    _font(signer.add_run(company), size=13, bold=True)

    line = document.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(line.add_run(f'{ceo_title} (인)'), size=13, bold=True)
    if seal:
        # 이름 옆에 겹쳐 찍지 않고 바로 아래에 둔다. 글자를 가리면 읽기 나빠진다.
        stamp = document.add_paragraph()
        stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        stamp.paragraph_format.space_before = Pt(2)
        stamp.add_run().add_picture(seal, width=Cm(2.2))

    document.save(path)
    return path
