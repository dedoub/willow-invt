#!/usr/bin/env python3
"""확정 급여대장 PDF 로 급여일 아침에 보낼 것 두 가지를 만든다.

  python3 scripts/tensw_payroll_payday.py 2026 8 8월급여대장.pdf \
      --register 급여내역_202608.xlsx --accounts 계좌.json --out 폴더

  1) 우리은행 대량이체 .xls — 머리글 없는 여덟 칸, 한 줄에 한 사람
  2) 개인별 급여명세서 — 워드(.docx, 고쳐 쓰라고)와 PDF(.pdf, 보내라고) 한 벌씩

이체 금액은 **차인지급액**이다. 우리가 다시 계산하지 않는다. 대장에 적힌 숫자를 옮긴다.
주민번호와 계좌번호는 깃·로그·위키에 적지 않는다. 계좌는 밖에서 받아 온 json 으로만 온다.
"""
import argparse
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl
import xlwt

sys.path.insert(0, str(Path(__file__).resolve().parent / 'lib'))
from kr_workdays import month_facts                                 # noqa: E402
from payroll_ledger import read_ledger                               # noqa: E402
from payslip_docx import build_payslip                               # noqa: E402
from payslip_pdf import build_payslip_pdf                            # noqa: E402

COMPANY = '텐소프트웍스'
COMPANY_FULL = '주식회사 텐소프트웍스'
CEO = '대표이사 김철형'

# 급여대장의 칸 → 명세서의 줄. 서식에 자리가 없는 것은 기타로 모은다.
PAY_ROWS = {'기본급': '기본급', '식대': '식대', '자가운전': '차량유지비'}
DEDUCT_ROWS = {
    '국민연금': '국민연금', '건강보험': '건강보험', '장기요양보험료': '장기요양보험',
    '고용보험': '고용보험', '소득세': '소득세', '지방소득세': '지방소득세',
}

DIGITS = '영일이삼사오육칠팔구'
SMALL_UNITS = ('', '십', '백', '천')
BIG_UNITS = ('', '만', '억', '조')


def korean_amount(value):
    """2,737,525 → 이백칠십삼만칠천오백이십오. 십 앞의 1 은 떼고 백·천 앞의 1 은 붙인다."""
    if value == 0:
        return '영'
    groups = []
    remainder = value
    while remainder:
        groups.append(remainder % 10_000)
        remainder //= 10_000
    parts = []
    for index in range(len(groups) - 1, -1, -1):
        group = groups[index]
        if not group:
            continue
        text = ''
        for place in range(3, -1, -1):
            digit = group // (10 ** place) % 10
            if not digit:
                continue
            text += ('' if digit == 1 and place == 1 else DIGITS[digit]) + SMALL_UNITS[place]
        parts.append(text + BIG_UNITS[index])
    return ''.join(parts)


def read_register(path):
    """급여내역 xlsx 에서 주민번호를 가져온다. 대장 PDF 에는 없다."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook.worksheets[0]
    header = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value}
    name_col, id_col = header.get('이름'), header.get('주민번호')
    if not name_col or not id_col:
        raise SystemExit(f'급여내역에 이름·주민번호 칸이 없어요: {list(header)}')
    out = {}
    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row, name_col).value
        if not name:
            continue
        out[str(name).strip()] = str(sheet.cell(row, id_col).value or '').strip()
    return out


def write_transfer(people, accounts, residents, year, month, destination):
    """우리은행 대량이체 — 은행, 계좌, 금액, 이름, 생년월일6자리, 빈칸, 보내는이, 적요.

    실제로 올라가는 파일은 여덟 칸이 모두 **글자**다. 금액을 숫자로 적으면 서식에 따라
    다르게 읽힌다. 계좌번호도 앞의 0 이 날아간다.
    """
    book = xlwt.Workbook(encoding='utf-8')
    sheet = book.add_sheet('Sheet1')
    note = f'{month}월급여'
    missing = []
    mismatched = []
    for index, person in enumerate(people):
        account = accounts.get(person['name'])
        if not account:
            missing.append(person['name'])
            account = {}
        # 생년월일은 계좌 장부에 적힌 것을 쓴다. 은행에 등록된 생년월일이 주민번호와
        # 다른 사람이 있어서, 주민번호에서 잘라 만들면 조용히 틀린다.
        resident = re.sub(r'\D', '', residents.get(person['name'], ''))[:6]
        birth = account.get('birth') or resident
        if account.get('birth') and resident and account['birth'] != resident:
            mismatched.append((person['name'], account['birth'], resident))
        cells = [
            account.get('bank', ''),
            account.get('number', ''),
            str(person['values']['차인지급액']),
            person['name'],
            birth,
            '',
            COMPANY,
            note,
        ]
        for column, value in enumerate(cells):
            sheet.write(index, column, value)
    book.save(destination)
    return missing, mismatched


def write_payslip(person, year, month, paid_on, destination, seal=None):
    """명세서 한 장을 워드로 쓴다. 돌려주는 것은 서식에 자리가 없어 기타로 간 금액이다."""
    values = person['values']

    pays, placed_pay = [], 0
    for key, label in PAY_ROWS.items():
        amount = values.get(key, 0)
        if amount:
            pays.append((label, amount))
            placed_pay += amount
    deducts, placed_deduct = [], 0
    for key, label in DEDUCT_ROWS.items():
        amount = values.get(key, 0)
        if amount:
            deducts.append((label, amount))
            placed_deduct += amount

    # 정산·두루누리·연말정산처럼 줄이 없는 항목을 조용히 버리면 합이 안 맞는다.
    pay_rest = values['지급합계'] - placed_pay
    deduct_rest = values['공제합계'] - placed_deduct
    if pay_rest:
        pays.append(('기타', pay_rest))
    if deduct_rest:
        deducts.append(('기타', deduct_rest))

    net = values['차인지급액']
    layout = dict(
        name=person['name'], department=person['department'], title=person['title'],
        period=f'{year}.{month:02d}.01~{year}.{month:02d}.{month_end(year, month):02d}',
        paid_on=paid_on,
        pays=pays, deducts=deducts,
        pay_total=values['지급합계'], deduct_total=values['공제합계'],
        net=net, net_korean=korean_amount(net),
        company=COMPANY_FULL, ceo_title=CEO, seal=seal,
    )
    # 워드는 고쳐 쓰라고 두고, 보낼 PDF 는 따로 그린다. 둘 다 같은 인자로 만든다.
    build_payslip(path=destination, **layout)
    build_payslip_pdf(path=destination.with_suffix('.pdf'), **layout)
    return pay_rest, deduct_rest


def month_end(year, month):
    import calendar
    return calendar.monthrange(year, month)[1]


ROSTER = {
    # 명세서의 부서·직책. 대장에는 직급만 있고 부서는 비어 있다.
    '김철형': ('경영총괄', '대표이사'),
    '김의향': ('경영지원', '이사'),
    '김정한': ('경영지원', '이사'),
    '박선영': ('경영지원', '이사'),
    '김경수': ('AI검색연구소', '선임 연구원'),
    '권지민': ('AI검색연구소', '연구원'),
    '조성민': ('AI검색연구소', '연구원'),
    '이승무': ('AI검색연구소', '연구원'),
    '전희나': ('AI검색연구소', '연구원'),
}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('year', type=int)
    parser.add_argument('month', type=int)
    parser.add_argument('ledger', help='세무법인이 보낸 확정 급여대장 PDF')
    parser.add_argument('--register', required=True, help='그 달 급여내역 xlsx (주민번호를 여기서 가져온다)')
    parser.add_argument('--accounts', help='{"이름": {"bank": "우리은행", "number": "1002…", "birth": "770818"}} 형태의 json')
    parser.add_argument('--seal', help='법인인감 png. 공식문서함에서 받아 온다:\n'
                        '  curl -s "$(npx tsx scripts/corp-records.ts doc url TS-DOC-2026-003 | tr -d \'"\')" -o /tmp/seal.png')
    parser.add_argument('--out', required=True)
    args = parser.parse_args()

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    people = read_ledger(args.ledger)
    residents = read_register(args.register)
    accounts = json.loads(Path(args.accounts).read_text()) if args.accounts else {}

    for person in people:
        department, title = ROSTER.get(person['name'], ('', person['rank']))
        person['department'] = department
        person['title'] = title

    paid_on = datetime.date.fromisoformat(month_facts(args.year, args.month)['payDate'])
    label = f'{args.year}{args.month:02d}'

    transfer = out / f'{COMPANY}_대량이체_{label}.xls'
    missing, mismatched = write_transfer(people, accounts, residents, args.year, args.month, transfer)

    print(f'{args.year}년 {args.month}월 · 지급일 {paid_on}')
    print(f'\n대량이체 → {transfer.name}  ({len(people)}명, 합계 '
          f'{sum(p["values"]["차인지급액"] for p in people):,}원)')
    if missing:
        print(f'  계좌를 모르는 사람: {", ".join(missing)}  ← 채워 넣어야 올릴 수 있어요')
    for name, birth, resident in mismatched:
        print(f'  {name}: 계좌 장부의 생년월일 {birth} 이 주민번호 앞자리 {resident} 와 달라요 (장부를 따랐어요)')

    print('\n급여명세서')
    for person in people:
        destination = out / f'급여명세서_{label}_{person["name"]}.docx'
        pay_rest, deduct_rest = write_payslip(person, args.year, args.month, paid_on, destination, args.seal)
        extra = []
        if pay_rest:
            extra.append(f'지급 기타 {pay_rest:,}')
        if deduct_rest:
            extra.append(f'공제 기타 {deduct_rest:,}')
        note = f'   ({", ".join(extra)})' if extra else ''
        print(f'  {person["name"]}  실수령 {person["values"]["차인지급액"]:,}원{note}')


if __name__ == '__main__':
    main()
