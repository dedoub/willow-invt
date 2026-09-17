#!/usr/bin/env python3
"""텐소프트웍스 급여내역(세무법인에 보낼 입력본)을 만든다.

  python3 scripts/tensw_payroll_register.py 2026 9 <직전달.xlsx> <출력.xlsx> [내려받은csv…]

세무법인에 보내는 것은 계산서가 아니라 **입력**이다. 4대보험 사이트가 준 숫자를 서식의
제자리에 옮겨 담을 뿐, 근로자 부담분을 여기서 계산하지 않는다. 확정 급여대장은 세무법인이 만든다.

사회보험 사이트 자료가 급여일 오전까지 안 올라오면 그 보험은 직전 달 숫자를 그대로 쓴다(CEO).
어느 칸을 새로 받았고 어느 칸을 이어썼는지 표로 찍어 준다 — 조용히 이어쓰면 틀려도 모른다.
"""
import csv, sys, shutil, io
from pathlib import Path
import openpyxl

HEADER_ROW = 1
NAME_COL = 2            # B: 이름

# 사회보험 사이트가 주는 칸 ← 급여대장의 칸
FROM_SITE = {
    '건강': '건강',
    '요양': '요양',
    '국민연금': '국민연금(합산)',
    '고용': '고용(합산)',
    '산재': '산재(합산)',
}

def read_nhis_csv(path):
    """건강보험공단이 주는 개인별 CSV. 앞쪽에 빈 줄이 붙어 오고 인코딩은 EUC-KR 이다."""
    raw = Path(path).read_bytes()
    for enc in ('euc-kr', 'cp949', 'utf-8-sig', 'utf-8'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SystemExit(f'인코딩을 못 읽었어요: {path}')
    lines = [l for l in text.splitlines() if l.strip()]
    head = next((i for i, l in enumerate(lines) if l.startswith('순번')), None)
    if head is None:
        raise SystemExit(f'머리글(순번)을 못 찾았어요: {path}')
    rows = list(csv.reader(lines[head:]))
    header, body = rows[0], rows[1:]
    # 같은 이름의 칸이 두 벌(건강/요양) 이어 붙는다. 구분 칸 위치로 갈라 읽는다.
    kinds = [i for i, h in enumerate(header) if h == '구분']
    out = {}
    for row in body:
        if len(row) < 4 or not row[3].strip():
            continue
        name = row[3].strip()
        per = out.setdefault(name, {})
        for start in kinds:
            kind = row[start].strip()
            if not kind:
                continue
            # 구분 다음 칸이 산출보험료
            per[kind] = int(float(row[start + 1] or 0))
    return out

def main():
    year, month = int(sys.argv[1]), int(sys.argv[2])
    previous, output = sys.argv[3], sys.argv[4]
    csvs = sys.argv[5:]

    fetched = {}
    for path in csvs:
        for name, values in read_nhis_csv(path).items():
            fetched.setdefault(name, {}).update(values)

    shutil.copyfile(previous, output)
    wb = openpyxl.load_workbook(output)
    ws = wb.worksheets[0]
    header = {str(c.value).strip(): c.column for c in ws[HEADER_ROW] if c.value}
    for want in FROM_SITE.values():
        if want not in header:
            raise SystemExit(f'서식에 "{want}" 칸이 없어요. 머리글: {list(header)}')

    report = []
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        name = ws.cell(row, NAME_COL).value
        if not name:
            continue
        name = str(name).strip()
        marks = []
        for site_key, column in FROM_SITE.items():
            cell = ws.cell(row, header[column])
            new = fetched.get(name, {}).get(site_key)
            if new is None:
                marks.append(f'{site_key}=이어씀')
            else:
                if cell.value != new:
                    marks.append(f'{site_key}={cell.value}→{new}')
                else:
                    marks.append(f'{site_key}=같음')
                cell.value = new
        report.append((name, marks))

    wb.save(output)
    print(f'{year}년 {month}월 급여내역 → {output}')
    print(f'내려받은 자료: {", ".join(Path(c).name for c in csvs) if csvs else "없음 (전부 직전 달 숫자)"}\n')
    for name, marks in report:
        print(f'  {name:6s} {"  ".join(marks)}')

if __name__ == '__main__':
    main()
